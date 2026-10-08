#!/usr/bin/env python3
"""Compara dos libros Excel hoja a hoja y celda a celda.

Compara valores y fórmulas tal como están guardados en el fichero (no recalcula).
"""
from __future__ import annotations

import argparse
import sys

import openpyxl
from openpyxl.chartsheet import Chartsheet


def _iguales(a, b, tolerancia: float) -> bool:
    # bool es subclase de int en Python: sin este chequeo, True == 1 se consideraría
    # «igual» con la comparación numérica de abajo, aunque en un Excel real TRUE y 1 son
    # valores de tipo distinto y no deberían darse por iguales.
    if isinstance(a, bool) or isinstance(b, bool):
        return isinstance(a, bool) and isinstance(b, bool) and a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= tolerancia * max(1.0, abs(a), abs(b))
    return a == b


def comparar(a, b, tolerancia: float = 1e-9, ignorar=()) -> list[str]:
    ignorar = {i.lower() for i in ignorar}
    la = openpyxl.load_workbook(a, data_only=False)
    lb = openpyxl.load_workbook(b, data_only=False)
    diferencias = [f"hoja solo en {n}: {h}" for n, libro, otro in (("A", la, lb), ("B", lb, la))
                   for h in libro.sheetnames if h not in otro.sheetnames]
    for hoja in [h for h in la.sheetnames if h in lb.sheetnames]:
        ha, hb = la[hoja], lb[hoja]
        # Las hojas de gráfico (chartsheet) no tienen celdas: compararlas como si fueran
        # una hoja normal provoca un error. Se avisa y se saltan en vez de fallar.
        es_grafico_a, es_grafico_b = isinstance(ha, Chartsheet), isinstance(hb, Chartsheet)
        if es_grafico_a or es_grafico_b:
            if es_grafico_a != es_grafico_b:
                diferencias.append(f"{hoja}: es hoja de gráfico en un libro y hoja normal en el otro")
            else:
                # Esto es un aviso, no una diferencia: no cuenta para el resultado (si no,
                # comparar un libro con gráfico contra sí mismo saldría con diferencias).
                print(f"aviso: {hoja}: hoja de gráfico omitida (no se compara su contenido)",
                      file=sys.stderr)
            continue
        for fila in range(1, max(ha.max_row, hb.max_row) + 1):
            for col in range(1, max(ha.max_column, hb.max_column) + 1):
                va, vb = ha.cell(fila, col).value, hb.cell(fila, col).value
                ref = f"{hoja}!{ha.cell(fila, col).coordinate}"
                if ref.lower() not in ignorar and not _iguales(va, vb, tolerancia):
                    diferencias.append(f"{ref}: {va!r} != {vb!r}")
        rangos_a = {str(r) for r in ha.merged_cells.ranges}
        rangos_b = {str(r) for r in hb.merged_cells.ranges}
        for n, rangos, otros in (("A", rangos_a, rangos_b), ("B", rangos_b, rangos_a)):
            for rango in sorted(rangos - otros):
                diferencias.append(f"{hoja}: celdas combinadas {rango} solo en {n}")
    return diferencias


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("a")
    p.add_argument("b")
    p.add_argument("--ignorar", nargs="*", default=[])
    p.add_argument("--tolerancia", type=float, default=1e-9)
    args = p.parse_args(argv)
    diferencias = comparar(args.a, args.b, args.tolerancia, args.ignorar)
    print(f"{len(diferencias)} diferencias")
    for d in diferencias[:200]:
        print("  " + d)
    return 0 if not diferencias else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    sys.exit(main())
